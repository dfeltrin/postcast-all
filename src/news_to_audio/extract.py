from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

import httpx
import trafilatura
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from pptx import Presentation

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".csv", ".docx", ".pptx"}


class SourceAccessError(ValueError):
    """The source cannot currently be retrieved by the application."""


@dataclass(frozen=True)
class ExtractedSource:
    title: str
    text: str
    content_type: str


def _title_from_url(url: str) -> str:
    path = urlparse(url).path.rstrip("/")
    return Path(path).stem or urlparse(url).netloc or "source"


def _clean(text: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _html_to_text(body: bytes, url: str) -> ExtractedSource:
    html = body.decode("utf-8", errors="replace")
    extracted = trafilatura.extract(html, include_comments=False, include_tables=True)
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else _title_from_url(url)
    fallback = soup.get_text("\n", strip=True)
    return ExtractedSource(title, _clean(extracted or fallback), "text/html")


def _file_to_text(body: bytes, extension: str, title: str) -> ExtractedSource:
    if extension == ".pdf":
        text = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(body)).pages)
    elif extension in {".txt", ".md"}:
        text = body.decode("utf-8", errors="replace")
    elif extension == ".csv":
        rows = csv.reader(io.StringIO(body.decode("utf-8", errors="replace")))
        text = "\n".join(" | ".join(row) for row in rows)
    elif extension == ".docx":
        document = Document(io.BytesIO(body))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    elif extension == ".pptx":
        presentation = Presentation(io.BytesIO(body))
        text = "\n\n".join(
            "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
            for slide in presentation.slides
        )
    else:
        raise ValueError(f"Formato non supportato: {extension or 'sconosciuto'}")
    return ExtractedSource(title, _clean(text), extension.lstrip("."))


def extract_url(url: str, timeout: float = 30.0) -> ExtractedSource:
    """Download a public URL and return meaningful textual content.

    Authentication and paywall bypassing are intentionally out of scope.
    """
    headers = {"User-Agent": "PostcastAll/0.1 (+local personal automation)"}
    try:
        with httpx.Client(follow_redirects=True, timeout=timeout, headers=headers) as client:
            response = client.get(url)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        status = exc.response.status_code if isinstance(exc, httpx.HTTPStatusError) else None
        detail = f"HTTP {status}" if status is not None else str(exc)
        raise SourceAccessError(f"Pagina non accessibile: {detail}") from exc
    content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
    extension = Path(urlparse(str(response.url)).path).suffix.lower()
    if content_type == "application/pdf":
        extension = ".pdf"
    if extension in SUPPORTED_EXTENSIONS:
        result = _file_to_text(response.content, extension, _title_from_url(str(response.url)))
    elif content_type in {"text/html", "application/xhtml+xml"} or not content_type:
        result = _html_to_text(response.content, str(response.url))
    else:
        raise ValueError(f"Tipo di contenuto non supportato: {content_type}")
    if not result.text:
        raise ValueError("Non e' stato possibile estrarre testo dalla fonte.")
    return result
